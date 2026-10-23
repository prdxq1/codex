"""Запуск: python tools/build_report.py [--name 'Фамилия И. О.'] [--group '...']."""
from pathlib import Path
from copy import deepcopy
import argparse,csv,json
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
a=argparse.ArgumentParser();a.add_argument('--name',default='Дутов Р. В.');a.add_argument('--group',default='');a.add_argument('--pages');args=a.parse_args()
d=Document(ROOT/'tools/report_template.docx')
# Переписанные пояснения заменяют текст, но сохраняют оформление абзацев.
content=json.loads((ROOT/'tools/report_text.json').read_text())
for index,text in content.items():
 p=d.paragraphs[int(index)]
 style=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear()
 run=p.add_run(text)
 if style is not None:run._r.insert(0,style)

# Оглавление образца содержит устаревшие разделы; сохраняем его стили,
# отступы и заполнитель, заменяя состав и номера страниц.
toc_titles=['1. Постановка задачи','2. Математическое описание алгоритма','2.1. Алгоритмы выделения контуров','2.2. Линейные усредняющие фильтры','3. Программное описание алгоритма','4. Результат работы программы','5. Заключение','Приложение А','Приложение Б']
pages=json.loads(Path(args.pages).read_text()) if args.pages else [3,4,4,8,10,11,15,16,19]
for sdt in list(d._element.body):
 if sdt.tag!=qn('w:sdt'):continue
 content=sdt.find(qn('w:sdtContent'))
 if content is None:continue
 old=content.findall(qn('w:p'))
 if not old:continue
 heading=deepcopy(old[0]);entry=deepcopy(old[1])
 for child in list(content):content.remove(child)
 def textpara(element,text):
  for child in list(element):
   if child.tag!=qn('w:pPr'):element.remove(child)
  para=Paragraph(element,d._body);para.add_run(text);return para
 hp=textpara(heading,'Содержание');hp.paragraph_format.page_break_before=True;content.append(heading)
 for title,page in zip(toc_titles,pages):
  node=deepcopy(entry);p=textpara(node,title+'\t'+str(page))
  p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  content.append(node)
for p in d.paragraphs:
 if p.text.startswith('Студент:'):
  p.runs[0].text='Студент:\t'+args.name+('  Группа: '+args.group if args.group else '')
  for r in p.runs[1:]:r.text=''
rows=list(csv.DictReader((ROOT/'test_output/metrics.csv').open()))
metrics={r['method']:r for r in rows}
for p in d.paragraphs:
 if '{{METRICS}}' in p.text:
  text='Результаты измерений относительно чистого оригинала (MSE; PSNR, дБ; MSE на однородном фоне): '
  labels={'noisy':'шум','box3':'окно 3×3','box5':'окно 5×5','gaussian':'Гаусс 5×5, σ=1,5'}
  text+='; '.join(f"{labels[r['method']]} — {float(r['mse']):.2f}; {float(r['psnr_db']):.2f}; {float(r['flat_mse']):.2f}" for r in rows)+'.'
  p.runs[0].text=text
  for r in p.runs[1:]:r.text=''
images=['test_input/input.png','test_output/var1_laplacian.png','test_output/var1_roberts.png','test_output/var1_sobel.png','test_output/var2_box_blur_3x3.png','test_output/var2_box_blur_5x5.png','test_output/var2_gaussian_blur.png']
for shape,path in zip(d.inline_shapes,images):
 blip=shape._inline.xpath('.//a:blip')[0]
 rid, _=d.part.get_or_add_image(str(ROOT/path));blip.set(qn('r:embed'),rid)
# Удаляем неиспользуемые связи с изображениями шаблона.
used={b.get(qn('r:embed')) for b in d._element.xpath('.//a:blip')}
for rid,rel in list(d.part.rels.items()):
 if rel.reltype.endswith('/image') and rid not in used:d.part.drop_rel(rid)
d.core_properties.author=args.name if '_' not in args.name else ''
d.core_properties.last_modified_by=d.core_properties.author
d.core_properties.title='Цифровая обработка изображений. Варианты 1 и 2'
d.save(ROOT/'report/Лабораторная работа.docx')
print('Saved report/Лабораторная работа.docx')
